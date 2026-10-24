"""Portfolio report: fundamental evidence and technical timing kept separate.
Broker CSV/XLSX parsing is reused. Caller owns request directory and deletes it.
"""
import io
import os
import re
import csv
import math
import zipfile
import logging
from datetime import datetime, timedelta, timezone
from typing import Callable, Dict, List, Optional, Tuple

from portfolio_analytics import positions, totals, exposure, fundamental_view, number, annual_cagr, consecutive_periods

logger = logging.getLogger(__name__)
IST = timezone(timedelta(hours=5, minutes=30))
MAX_ROWS = 50
MAX_FILE_BYTES = 1_000_000
TEMPLATE_CSV = "Symbol,Qty,Buy Price\nRELIANCE,50,1250\nSBIN,100,750\n"

def _norm_header(h) -> str:
    return re.sub(r"[^a-z0-9 ]", "", str(h or "").lower()).strip()


_SYMBOL_H = {"symbol", "scrip", "scrip name", "scripname", "instrument", "stock", "stock name", "stockname",
             "tradingsymbol", "trading symbol", "ticker", "security", "security name", "company", "company name", "name"}
_QTY_H = {"qty", "quantity", "shares", "no of shares", "units", "holding qty", "holdings", "net qty", "balance qty"}
_PRICE_H = {"buy price", "buyprice", "avg price", "average price", "avg cost", "average cost", "avg buy price",
            "average buy price", "buy rate", "buy avg", "avgprice", "purchase price", "cost price", "rate", "price", "cost",
            "avg buy", "buy average"}


def _to_float(x) -> Optional[float]:
    if x is None:
        return None
    t = str(x).strip()
    if not t:
        return None
    neg = t.startswith("(") and t.endswith(")")
    t = re.sub(r"(?i)(₹|rs\.?|inr)", "", t).replace(",", "").replace("(", "").replace(")", "").strip()
    try:
        v = float(t)
    except ValueError:
        return None
    return (-v if neg else v) if math.isfinite(v) else None


def clean_symbol(raw, resolver: Optional[Callable[[str], Optional[str]]] = None) -> Optional[str]:
    s = str(raw or "").strip().upper()
    if not s or s.startswith("BSE:") or s.endswith((".BO", ".BSE")):
        return None
    s = re.sub(r"^(NSE|BSE|NSE_EQ)\s*[:\-]\s*", "", s)
    s = re.sub(r"\.(NS|BO|NSE|BSE)$", "", s)
    s = re.sub(r"-(EQ|BE|SM|ST)$", "", s)
    s = s.strip()
    if re.fullmatch(r"[A-Z0-9&\-]{1,20}", s):
        return s
    if resolver and re.search(r"[A-Za-z]{3,}", s):          # a company name like "Reliance Industries"
        try:
            r = resolver(str(raw).strip())
            return r.upper() if r else None
        except Exception:
            return None
    return None


def _merge(holdings: List[dict]) -> List[dict]:
    """Duplicate symbols (several lots) → one row with the weighted-average buy price."""
    agg: Dict[str, dict] = {}
    for h in holdings:
        a = agg.setdefault(h["symbol"], {"symbol": h["symbol"], "qty": 0.0, "cost": 0.0})
        a["qty"] += h["qty"]
        a["cost"] += h["qty"] * h["buy"]
    return [{"symbol": a["symbol"], "qty": a["qty"], "buy": (a["cost"] / a["qty"]) if a["qty"] else 0.0}
            for a in agg.values()]


def parse_rows(rows: List[List], resolver=None) -> Tuple[List[dict], List[str]]:
    """rows: list of cell lists (strings or numbers). → (holdings, issues)."""
    issues: List[str] = []
    rows = [r for r in rows if r and any(str(c or "").strip() for c in r)]
    if not rows:
        return [], ["The file is empty."]

    col = {"symbol": None, "qty": None, "price": None}
    header_at = None
    for i, r in enumerate(rows[:6]):
        norm = [_norm_header(c) for c in r]
        found = {"symbol": None, "qty": None, "price": None}
        for j, c in enumerate(norm):
            if c in _SYMBOL_H and found["symbol"] is None:
                found["symbol"] = j
            elif c in _QTY_H and found["qty"] is None:
                found["qty"] = j
            elif c in _PRICE_H and found["price"] is None:
                found["price"] = j
        if sum(v is not None for v in found.values()) >= 2:
            col, header_at = found, i
            break
    if header_at is None:                                     # no header → positional Symbol, Qty, Price
        col = {"symbol": 0, "qty": 1, "price": 2}
    if None in col.values():
        missing = [k for k, v in col.items() if v is None]
        return [], [f"Could not find the {', '.join(missing)} column. Expected headers like: Symbol, Qty, Buy Price."]

    out: List[dict] = []
    start = (header_at + 1) if header_at is not None else 0
    for n, r in enumerate(rows[start:], start=start + 1):
        if len(out) >= MAX_ROWS:
            return [], [f"More than {MAX_ROWS} holdings. Split the upload; no rows were silently omitted."]
        get = lambda k: r[col[k]] if col[k] is not None and col[k] < len(r) else None
        raw_sym = get("symbol")
        if re.match(r"(?i)^\s*(total|grand total|net)\b", str(raw_sym or "")):
            continue
        sym = clean_symbol(raw_sym, resolver)
        qty, buy = _to_float(get("qty")), _to_float(get("price"))
        if sym is None:
            issues.append(f"Row {n}: could not read a stock symbol from “{str(raw_sym)[:25]}”.")
        elif qty is None or qty <= 0:
            issues.append(f"Row {n} ({sym}): invalid quantity.")
        elif buy is None or buy < 0:
            issues.append(f"Row {n} ({sym}): invalid buy price.")
        else:
            out.append({"symbol": sym, "qty": qty, "buy": buy})
    merged = _merge(out)
    if len(merged) < len(out):
        issues.append(f"{len(out) - len(merged)} duplicate symbol row(s) were merged using the weighted-average buy price.")
    return merged, issues


def parse_text(text: str, resolver=None) -> Tuple[List[dict], List[str]]:
    rows = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if re.search(r"[,\t;|]", line):
            cells = [c.strip() for c in re.split(r"[,\t;|]", line)]
        else:
            cells = [c for c in re.split(r"\s+[xX]\s+|\s*[@×]\s*|\s+", line) if c]
        rows.append(cells)
    return parse_rows(rows, resolver)


def parse_file(name: str, data: bytes, resolver=None) -> Tuple[List[dict], List[str]]:
    if len(data) > MAX_FILE_BYTES:
        return [], ["File is too large (limit 1 MB)."]
    low = (name or "").lower()
    try:
        if low.endswith((".xlsx", ".xlsm")):
            import openpyxl
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 10_000_000:
                    return [], ["Expanded workbook exceeds 10 MB; export a smaller CSV."]
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            if ws.max_row and ws.max_row > MAX_ROWS + 6:
                wb.close()
                return [], [f"Workbook exceeds {MAX_ROWS} holdings plus header rows."]
            rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(max_row=MAX_ROWS + 10, max_col=30, values_only=True)]
            wb.close()
            return parse_rows(rows, resolver)
        if low.endswith(".xls"):
            return [], ["Old .xls files are not supported — please save as .xlsx or .csv."]
        text = None
        for enc in ("utf-8-sig", "utf-16", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except Exception:
                continue
        if text is None:
            return [], ["Could not read the file encoding."]
        try:
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t|")
            rows = list(csv.reader(io.StringIO(text), dialect))
        except Exception:
            rows = list(csv.reader(io.StringIO(text)))
        if rows and max(len(r) for r in rows) < 2:           # plain "SYMBOL QTY PRICE" lines
            return parse_text(text, resolver)
        return parse_rows(rows, resolver)
    except Exception as e:
        logger.warning(f"[portfolio] file parse failed: {type(e).__name__}")
        return [], ["Could not read the file. Please send a CSV or .xlsx with columns Symbol, Qty, Buy Price."]



def price_holdings(holdings):
    from market_data import bulk_history, index_history
    from market_breadth import completed_prices
    # One common completed session, so stale symbols cannot silently enter totals.
    bench = completed_prices(index_history('^NSEI', '1mo'))
    if bench.empty:
        return positions(holdings, {}), {}
    session = bench.index[-1]
    hist = bulk_history([h['symbol'] for h in holdings], '1y', cache_ttl=300)
    prices, clean = {}, {}
    for sym, frame in hist.items():
        df = completed_prices(frame)
        clean[sym] = df
        if not df.empty and session in df.index:
            px = number(df.loc[session, 'Close'])
            if px is not None and px > 0:
                prices[sym] = {'price': px, 'date': str(session.date())}
    return positions(holdings, prices), clean


def _analyse(sym, frame):
    from fundamentals import get_fundamentals
    from research_data import get_statements, get_screener_extras
    from research_engine import compute_fundamentals
    from nifty500_collector import SECTOR_STOCKS
    fund = get_fundamentals(sym) or {}
    stm = get_statements(sym) or {}
    extras = get_screener_extras(sym) or {}
    sector = next((sec for sec, syms in SECTOR_STOCKS.items() if sym in syms), 'Unclassified')
    lender = any(word in sector.lower() for word in ('bank', 'financ', 'nbfc', 'insurance'))
    px = float(frame['Close'].iloc[-1]) if frame is not None and not frame.empty else 0
    fv = compute_fundamentals(fund, stm, extras, [], 'Financial' if lender else sector, px)
    years, quarters = stm.get('years') or [], stm.get('quarters') or []
    # Never prefer a vendor CAGR over the annual table displayed to the client.
    for span in (3,5):
        for prefix,key in (('rev','revenue'),('pat','net_profit')):
            fv[f'{prefix}_cagr{span}'] = annual_cagr(years,key,span)
    if len(years) < 3 or not consecutive_periods(years[-3:],12):
        fv['ocf_pat'] = None
    if len(quarters) < 5 or not consecutive_periods(quarters[-5:],3):
        fv['q_pat_yoy'] = fv['q_rev_yoy'] = None
    review = fundamental_view(fv, stm, today=datetime.now(IST).date())
    if sector == 'Unclassified':
        review['view'] = 'Insufficient verified information'
        review['gaps'].append('Industry classification unavailable; sector-specific framework unverified')
    review.update(fv=fv, statements=stm, technical='Insufficient completed price history')
    if frame is not None and len(frame) >= 200:
        c = frame['Close']
        if c.iloc[-1] > c.iloc[-50:].mean() > c.iloc[-200:].mean():
            review['technical'] = 'Above 50/200-day SMA; positive price structure'
        elif c.iloc[-1] < c.iloc[-50:].mean() < c.iloc[-200:].mean():
            review['technical'] = 'Below 50/200-day SMA; weak price structure'
        else:
            review['technical'] = 'Mixed price structure'
    return review, sector


def build_summary_pdf(rows, path, client='', notes=()):
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak
    from config import BRAND_NAME
    styles = getSampleStyleSheet()
    # Embed ReportLab's bundled font, independent of host font substitutions.
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_dir = os.path.join(os.path.dirname(reportlab.__file__), 'fonts')
    if 'PortfolioSans' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('PortfolioSans', os.path.join(font_dir, 'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('PortfolioBold', os.path.join(font_dir, 'VeraBd.ttf')))
    for key in ('BodyText','Title','Heading1','Heading2'):
        styles[key].fontName = 'PortfolioSans' if key == 'BodyText' else 'PortfolioBold'
    styles['BodyText'].fontSize = 8.5
    styles['BodyText'].leading = 11
    styles['BodyText'].spaceAfter = 3
    styles['Heading2'].fontSize = 11
    styles['Heading2'].leading = 14
    styles['Heading2'].spaceBefore = 5
    styles['Heading2'].spaceAfter = 4
    def P(text, style='BodyText'):
        return Paragraph(escape(str(text)), styles[style])
    story = [P(BRAND_NAME, 'Title'), P('Portfolio Fundamental Review', 'Heading1'),
             P(datetime.now(IST).strftime('%d-%b-%Y %H:%M IST'))]
    if client: story.append(P(client))
    def section(title):
        story.extend([Spacer(1, 5), P(title, 'Heading2')])
    def table(headers, data, widths):
        t = LongTable([[P(x) for x in headers]] + [[P(x) for x in r] for r in data],
                      colWidths=[x*mm for x in widths], repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE8F4')),
                               ('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),4),
                               ('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#203A55'))]))
        story.append(t)
    fmt = lambda x, suffix='': 'Unavailable' if number(x) is None else f'{float(x):,.2f}{suffix}'
    t, e = totals(rows), exposure(rows)
    section('1. Portfolio value and evidence coverage')
    label = 'Total' if t['complete'] else 'PRICED SUBSET ONLY'
    story.append(P(f"{label}: Rs {t['value']:,.2f}; cost Rs {t['cost']:,.2f}; unrealised P&L Rs {t['pnl']:,.2f}."))
    story.append(P(f"Prices available: {t['priced']}/{t['total']} holdings. Fundamental review attempted for {e['analysed_pct']:.1f}% of priced value; this includes incomplete assessments."))
    if not t['complete']: story.append(P('Missing prices are excluded, never valued at purchase cost. Total portfolio value and P&L are unknown. All exposure percentages refer only to priced holdings.'))
    dates = sorted({r['price_date'] for r in rows if r['priced']})
    story.append(P('Price session: ' + ', '.join(dates) if dates else 'No verified completed-session prices.'))
    table(['Holding','Value (Rs)','Weight','P&L (Rs)'],
          [[r['symbol'],fmt(r['value']),fmt(r['weight_pct'],'%'),fmt(r['pnl'])] for r in rows], [40,45,35,50])
    section('2. Fundamental exposure and concentration')
    table(['Assessment','% of priced value'], [[k,fmt(v,'%')] for k,v in e['buckets'].items()], [120,50])
    table(['Sector / classification','% of priced value'], [[k,fmt(v,'%')] for k,v in sorted(e['sectors'].items(),key=lambda x:-x[1])], [120,50])
    for r in rows:
        if r['weight_pct'] is not None and r['weight_pct'] > 15:
            story.append(P(f"Concentration review: {r['symbol']} is {r['weight_pct']:.1f}% of priced value (15% review trigger, not a target allocation)."))
    for sec, weight in e['sectors'].items():
        if weight > 35: story.append(P(f'Sector concentration review: {sec} {weight:.1f}% (35% review trigger).'))
    for r in rows:
        if not r.get('review'):
            story.append(P(f"{r['symbol']}: fundamental analysis unavailable; no conclusion."))
    for i,r in enumerate(rows,1):
        if not r.get('review'):
            continue
        story.append(PageBreak())
        section(f"3.{i} {r['symbol']} - investment evidence")
        review = r.get('review')
        if not review:
            story.append(P('Analysis unavailable. No investment conclusion.'))
            continue
        fv = review['fv']
        story.append(P('Fundamental view: ' + review['view']))
        story.append(P('Valuation: ' + review['valuation']))
        story.append(P('Technical timing (separate): ' + review['technical']))
        story.append(P(f"Financial source: {review['source']}. Annual period: {review['annual']}; quarter: {review['quarter']}."))
        metrics = [('ROE',fv.get('roe'),'%'),('Revenue CAGR (3Y)',fv.get('rev_cagr3'),'%'),
                   ('Profit CAGR (3Y)',fv.get('pat_cagr3'),'%'),('Quarter profit YoY',fv.get('q_pat_yoy'),'%'),
                   ('P/E',fv.get('pe'),'x'),('P/B',fv.get('pb'),'x')]
        if not review['lender']:
            metrics += [('ROCE',fv.get('roce'),'%'),('Debt/equity',fv.get('de'),'x'),
                        ('Cash conversion: 3Y OCF/PAT',fv.get('ocf_pat'),'x'),('Operating margin change',fv.get('opm_delta'),' pp')]
        table(['Measure','Observed'], [[name,fmt(v,suffix)] for name,v,suffix in metrics], [120,50])
        for title,key in [('Supporting evidence','strengths'),('Risks / counter-evidence','risks'),('Missing evidence','gaps')]:
            section(title)
            for text in review[key] or ['None established from the available measures.']:
                story.append(P('- ' + text))
        annual_rows = (review.get('statements') or {}).get('years') or []
        if annual_rows:
            section('Annual earnings and cash flow (Rs crore)')
            story.append(P('CAGR uses dated annual endpoints; nonconsecutive periods or nonpositive endpoints produce no CAGR.'))
            table(['Period','Revenue','Net profit','Operating cash flow'],
                  [[y.get('label','Unknown'),fmt(y.get('revenue')),fmt(y.get('net_profit')),fmt(y.get('ocf')) if not review['lender'] else 'Not comparable']
                   for y in annual_rows[-4:]], [32,40,40,58])
        section('Next review / what changes the view')
        story.append(P(review['next_review']))
        story.append(P('Changes in asset quality, credit costs, profitability and capital adequacy require a fresh lender review.' if review['lender'] else 'An earnings decline, weaker cash conversion, rising debt or a materially richer valuation warrants a fresh review. A price trend change affects technical timing only.'))
    story.append(PageBreak()); section('Method, limitations and sources')
    for text in notes: story.append(P(text))
    for text in [
        'Prices: Yahoo Finance daily completed-session data. Financials: existing Screener.in / Yahoo / Finnhub fallback chain; exact financial source and period shown where supplied.',
        'Generic descriptive thresholds are review aids, not calibrated return forecasts. Sound fundamentals require current annual and quarterly evidence, all required measures, at least three supporting tests and no flagged concern.',
        'Lenders: industrial debt/equity, ROCE and operating cash conversion are not used for their conclusion. Asset quality, capital adequacy, ROA and credit costs are not verified by this feed; lender conclusions remain incomplete.',
        'No promoter pledge, auditor qualification, management guidance or corporate-event claim is inferred when verified information is absent. Insurance and other financial businesses may need an additional specialist framework.',
        'Peer multiples and a weighted portfolio P/E are not fabricated. A low P/E alone is not evidence of undervaluation. Exposure uses current market value; unknown prices remain unknown.',
        'No historical performance, personalised target weights, automatic buy/sell instruction or tax-aware recommendation is produced. Educational review; not SEBI-registered investment advice.'
    ]: story.append(P(text))
    def footer(canvas, doc):
        canvas.setFont('PortfolioSans',8); canvas.drawString(20*mm,12*mm,f'Portfolio review | Page {doc.page}')
    SimpleDocTemplate(path,pagesize=A4,leftMargin=20*mm,rightMargin=20*mm,
                      topMargin=18*mm,bottomMargin=20*mm).build(story,onFirstPage=footer,onLaterPages=footer)
    return path


def run_job(holdings, out_dir, client='', progress=None, max_analyze=None, use_ai=False):
    if not holdings or len(holdings) > MAX_ROWS:
        raise ValueError(f'Provide 1-{MAX_ROWS} holdings per review.')
    os.makedirs(out_dir, exist_ok=True)
    say = progress or (lambda text: None)
    rows, hist = price_holdings(holdings)
    notes = []
    # Every submitted holding gets a basic fundamental attempt; no top-10 blind spot.
    for i,r in enumerate(rows,1):
        say(f"Reviewing {r['symbol']} ({i}/{len(rows)})...")
        try:
            r['review'], r['sector'] = _analyse(r['symbol'], hist.get(r['symbol']))
        except Exception as exc:
            logger.warning('Portfolio fundamental review failed: %s',type(exc).__name__)
            notes.append(f"Holding {i}: fundamental analysis unavailable; no conclusion assigned.")
    path = os.path.join(out_dir,'Portfolio_Fundamental_Review.pdf')
    build_summary_pdf(rows,path,client,notes)
    return {'summary':path,'stocks':[],'notes':notes,'recs':rows}
